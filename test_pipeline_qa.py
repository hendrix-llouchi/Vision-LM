"""
QA & Pipeline Verification Script for Vision-LM
"""
import os
import sys
import json
import base64
import difflib
import io
from pathlib import Path
from PIL import Image
import openpyxl
import pandas as pd

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT_ROOT))

import pipeline
import app
from test_fuzzy import is_duplicate

def run_tests():
    print("=" * 70)
    print("VISION-LM QA & PIPELINE TEST SUITE")
    print("=" * 70)
    
    test_results = {}
    
    # -------------------------------------------------------------
    # STAGE 1: IMAGE INGESTION & PREPROCESSING
    # -------------------------------------------------------------
    print("\n--- [TEST 1] STAGE 1: IMAGE INGESTION & PREPROCESSING ---")
    sample_dir = PROJECT_ROOT / "sample_images"
    assert sample_dir.exists(), f"Sample directory missing: {sample_dir}"
    
    # Test pipeline.load_and_group_images
    groups = pipeline.load_and_group_images(sample_dir)
    print(f"Discovered groups: {len(groups)} products, {sum(len(v) for v in groups.values())} total images")
    
    assert len(groups) == 10, f"Expected 10 product groups, found {len(groups)}"
    total_imgs = sum(len(v) for v in groups.values())
    assert total_imgs == 10, f"Expected 10 total images, found {total_imgs}"
    
    # Verify each image can be preprocessed with both pipeline.py and app.py
    image_details = []
    for pid, paths in groups.items():
        for p in paths:
            # Check file readability and PIL
            with Image.open(p) as img:
                orig_w, orig_h = img.size
                orig_mode = img.mode
                orig_format = img.format
            
            # Test pipeline.preprocess_image
            b64_pipe = pipeline.preprocess_image(p)
            assert isinstance(b64_pipe, str) and len(b64_pipe) > 0, f"Pipeline b64 failed for {p.name}"
            
            # Decode and verify dimensions <= 1024x1024
            img_bytes = base64.b64decode(b64_pipe)
            with Image.open(io.BytesIO(img_bytes)) as decoded_img:
                dec_w, dec_h = decoded_img.size
                assert dec_w <= 1024 and dec_h <= 1024, f"Image {p.name} exceeded 1024x1024: ({dec_w}, {dec_h})"
                assert decoded_img.format == "JPEG", f"Image {p.name} format is not JPEG: {decoded_img.format}"
            
            # Test app.preprocess_image (using BytesIO / file-like)
            with open(p, "rb") as f:
                b64_app = app.preprocess_image(f)
            assert isinstance(b64_app, str) and len(b64_app) > 0, f"App b64 failed for {p.name}"
            
            image_details.append({
                "product_id": pid,
                "file_name": p.name,
                "orig_size": f"{orig_w}x{orig_h}",
                "orig_mode": orig_mode,
                "orig_format": orig_format,
                "processed_size": f"{dec_w}x{dec_h}",
                "b64_len": len(b64_pipe),
                "valid": True
            })
            print(f"  [OK] {p.name}: {orig_w}x{orig_h} ({orig_mode}) -> {dec_w}x{dec_h} (JPEG, {len(b64_pipe)} chars b64)")

    test_results["stage_1_image_ingestion"] = {
        "status": "PASSED",
        "total_images": len(image_details),
        "total_products": len(groups),
        "images": image_details
    }

    # -------------------------------------------------------------
    # STAGE 3 & 4: DEDUPLICATION, AGGREGATION, NORMALIZATION & 13 COLS
    # -------------------------------------------------------------
    print("\n--- [TEST 2] DEDUPLICATION LOGIC (test_fuzzy.py) ---")
    fuzzy_tests = [
        # Test 1: Near identical barcode (ratio > 0.85) -> duplicate
        ({"barcode": "599886628377", "brand": "KILIF", "item_name": "KILIF POWDER"},
         {"barcode": "599886628378", "brand": "KILIF", "item_name": "KILIF POWDER"}, True),
        # Test 2: Different barcode (< 0.85) -> not duplicate
        ({"barcode": "123456789012", "brand": "KILIF", "item_name": "KILIF POWDER"},
         {"barcode": "987654321098", "brand": "KILIF", "item_name": "KILIF POWDER"}, False),
        # Test 3: Same brand, same item name, weight mismatch -> not duplicate
        ({"item_name": "KIVO TOMATO MIX 60G", "barcode": "EMPTY", "brand": "KIVO", "weight": "60G"},
         {"item_name": "KIVO TOMATO MIX 100G", "barcode": "EMPTY", "brand": "KIVO", "weight": "100G"}, False),
        # Test 4: Similar item name (> 0.8), same brand, no weight conflict -> duplicate
        ({"item_name": "KILIF KONCENTRAT DETERGENT - POW", "barcode": "599886628377", "brand": "KILIF"},
         {"item_name": "KILIF CONCENTRATED DETERGENT PO", "barcode": "599886628377", "brand": "KIILIF"}, True),
        # Test 5: Different brands (> 0.8 ratio fails) -> not duplicate
        ({"item_name": "CHICKEN STOCK CUBE", "brand": "KNORR", "barcode": ""},
         {"item_name": "CHICKEN STOCK CUBE", "brand": "MAGGI", "barcode": ""}, False),
    ]
    for idx, (item_a, item_b, expected) in enumerate(fuzzy_tests, 1):
        res = is_duplicate(item_a, item_b)
        assert res == expected, f"Fuzzy test {idx} failed: got {res}, expected {expected}"
        print(f"  [OK] Fuzzy test {idx}: is_duplicate = {res} (matches expected: {expected})")
    
    test_results["deduplication_fuzzy"] = {"status": "PASSED", "tests_run": len(fuzzy_tests)}

    print("\n--- [TEST 3] AGGREGATION & CONFLICT RESOLUTION ---")
    multi_records = [
        {
            "ITEM_NAME": "MAGGI 2-MIN NOODLES CURRY 5X79G",
            "BARCODE": "9556001123456",
            "BRAND": "MAGGI",
            "MANUFACTURER": "NESTLE MALAYSIA",
            "WEIGHT": "5X79G",
            "PACKAGING_TYPE": "PACKET",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CURRY",
            "TYPE": "INSTANT NOODLES",
            "FRAGRANCE_FLAVOR": "CURRY",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": "2 MINIT"
        },
        {
            "ITEM_NAME": "MAGGI 2-MIN NOODLES CURRY 5X79G PACK", # Longer tie-breaker
            "BARCODE": "9556001123456",
            "BRAND": "MAGGI",
            "MANUFACTURER": "NESTLE PRODUCTS SDN BHD", # Longer tie-breaker
            "WEIGHT": "5X79G",
            "PACKAGING_TYPE": "PACKET",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CURRY",
            "TYPE": "INSTANT NOODLES",
            "FRAGRANCE_FLAVOR": "CURRY",
            "PROMOTION": "BUY 1 FREE 1",
            "ADDONS": "FREE BOWL",
            "TAGLINE": "2 MINIT"
        }
    ]
    agg_pipe = pipeline.aggregate(multi_records)
    agg_app = app.aggregate(multi_records)
    assert agg_pipe["PROMOTION"] == "BUY 1 FREE 1", "Aggregation failed to fill promotion"
    assert agg_pipe["ADDONS"] == "FREE BOWL", "Aggregation failed to fill addons"
    assert agg_pipe["BRAND"] == "MAGGI", "Aggregation failed brand majority"
    assert agg_pipe == agg_app, "pipeline.aggregate and app.aggregate outputs differ!"
    print(f"  [OK] Aggregation resolved correctly: {agg_pipe['ITEM_NAME']} | {agg_pipe['MANUFACTURER']}")
    test_results["aggregation"] = {"status": "PASSED"}

    print("\n--- [TEST 4] VALIDATION & NORMALIZATION ---")
    raw_record = {
        "ITEM_NAME": "  nestle Milo chocolate malt powder 1kg  ",
        "BARCODE": "955-6001-23456-7 ",
        "MANUFACTURER": " nestle manufacturing (malaysia) sdn bhd ",
        "BRAND": " milo ",
        "WEIGHT": " 1 kg ",
        "PACKAGING_TYPE": " pouch ",
        "COUNTRY": " malaysia ",
        "VARIANT": " active-go ",
        "TYPE": " malt drink ",
        "FRAGRANCE_FLAVOR": " chocolate ",
        "PROMOTION": " 10% extra ",
        "ADDONS": " none ",
        "TAGLINE": " energy to go further "
    }
    val_pipe = pipeline.validate_and_normalize(raw_record.copy())
    val_app = app.validate(raw_record.copy())
    
    assert val_pipe["BARCODE"] == "9556001234567", f"Barcode not cleaned: {val_pipe['BARCODE']}"
    assert val_pipe["WEIGHT"] == "1KG", f"Weight not normalized: {val_pipe['WEIGHT']}"
    assert val_pipe["ITEM_NAME"] == "NESTLE MILO CHOCOLATE MALT POWDER 1KG", f"Uppercase failed: {val_pipe['ITEM_NAME']}"
    assert val_pipe["BRAND"] == "MILO"
    assert val_pipe["PACKAGING_TYPE"] == "POUCH"
    assert val_pipe == val_app, "pipeline.validate_and_normalize and app.validate outputs differ!"
    print(f"  [OK] Barcode: '955-6001-23456-7 ' -> '{val_pipe['BARCODE']}'")
    print(f"  [OK] Weight: ' 1 kg ' -> '{val_pipe['WEIGHT']}'")
    print(f"  [OK] Uppercase: '{val_pipe['ITEM_NAME']}'")
    test_results["validation_normalization"] = {"status": "PASSED"}

    print("\n--- [TEST 5] 13 IMDB COLUMNS STRUCTURE ---")
    expected_cols = [
        "ITEM NAME", "BARCODE", "MANUFACTURER", "BRAND", "WEIGHT",
        "PACKAGING TYPE", "COUNTRY", "VARIANT", "TYPE",
        "FRAGRANCE FLAVOR", "PROMOTION", "ADDONS", "TAGLINE"
    ]
    assert pipeline.IMDB_COLS == expected_cols, "pipeline.IMDB_COLS does not match 13 hackathon columns"
    imdb_row = pipeline.to_imdb_row(val_pipe)
    assert list(imdb_row.keys()) == expected_cols, "to_imdb_row columns order/names mismatch"
    print(f"  [OK] 13 Columns exact match: {list(imdb_row.keys())}")
    test_results["imdb_columns"] = {"status": "PASSED", "columns": expected_cols}

    # -------------------------------------------------------------
    # SIMULATED END-TO-END RUN ON 10 SAMPLE IMAGES
    # -------------------------------------------------------------
    print("\n--- [TEST 6] SIMULATED END-TO-END PIPELINE RUN ON 10 SAMPLE IMAGES ---")
    mock_catalog = {
        "S221234199": {
            "ITEM_NAME": "KNORR CHICKEN STOCK CUBE 60G",
            "BARCODE": "8850144201015",
            "MANUFACTURER": "UNILEVER (MALAYSIA) HOLDINGS SDN BHD",
            "BRAND": "KNORR",
            "WEIGHT": "60G",
            "PACKAGING_TYPE": "BOX",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CHICKEN",
            "TYPE": "STOCK CUBE / SEASONING",
            "FRAGRANCE_FLAVOR": "CHICKEN",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": "REAL CHICKEN TASTE"
        },
        "S221712802": {
            "ITEM_NAME": "SUNLIGHT DISHWASHING LIQUID LEMON 1L",
            "BARCODE": "8850144302021",
            "MANUFACTURER": "UNILEVER THAILAND",
            "BRAND": "SUNLIGHT",
            "WEIGHT": "1L",
            "PACKAGING_TYPE": "BOTTLE",
            "COUNTRY": "THAILAND",
            "VARIANT": "LEMON 100",
            "TYPE": "DISHWASHING LIQUID",
            "FRAGRANCE_FLAVOR": "LEMON",
            "PROMOTION": "BUY 2 FREE 1",
            "ADDONS": "",
            "TAGLINE": "FASTER DEGREASING POWER"
        },
        "S222775012": {
            "ITEM_NAME": "MAGGI 2-MINUTE NOODLES CURRY FLAVOUR 5X79G",
            "BARCODE": "9556001002232",
            "MANUFACTURER": "NESTLE PRODUCTS SDN BHD",
            "BRAND": "MAGGI",
            "WEIGHT": "395G",
            "PACKAGING_TYPE": "PACKET",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CURRY",
            "TYPE": "INSTANT NOODLES",
            "FRAGRANCE_FLAVOR": "CURRY",
            "PROMOTION": "VALUE PACK 5 PACKS",
            "ADDONS": "",
            "TAGLINE": "KEENAKAN RASA KARI SEBENAR"
        },
        "S222894050": {
            "ITEM_NAME": "COLGATE TOTAL CLEAN MINT TOOTHPASTE 150G",
            "BARCODE": "8850006341201",
            "MANUFACTURER": "COLGATE-PALMOLIVE (M) SDN BHD",
            "BRAND": "COLGATE",
            "WEIGHT": "150G",
            "PACKAGING_TYPE": "TUBE",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CLEAN MINT",
            "TYPE": "TOOTHPASTE",
            "FRAGRANCE_FLAVOR": "MINT",
            "PROMOTION": "FREE EXTRA 20G",
            "ADDONS": "",
            "TAGLINE": "12 HOUR ANTIBACTERIAL PROTECTION"
        },
        "S222985766": {
            "ITEM_NAME": "DUTCH LADY FULL CREAM MILK 1L",
            "BARCODE": "9556028120018",
            "MANUFACTURER": "DUTCH LADY MILK INDUSTRIES BERHAD",
            "BRAND": "DUTCH LADY",
            "WEIGHT": "1L",
            "PACKAGING_TYPE": "CARTON",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "FULL CREAM",
            "TYPE": "UHT MILK",
            "FRAGRANCE_FLAVOR": "CREAMY MILK",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": "STRONG NUTRITION FOR FAMILY"
        },
        "S225637028": {
            "ITEM_NAME": "MILO ACTIV-GO CHOCOLATE MALT POWDER REFILL 1KG",
            "BARCODE": "9556001221109",
            "MANUFACTURER": "NESTLE MANUFACTURING (MALAYSIA) SDN BHD",
            "BRAND": "MILO",
            "WEIGHT": "1KG",
            "PACKAGING_TYPE": "POUCH",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "ACTIV-GO",
            "TYPE": "CHOCOLATE MALT BEVERAGE",
            "FRAGRANCE_FLAVOR": "CHOCOLATE MALT",
            "PROMOTION": "SAVING PACK",
            "ADDONS": "",
            "TAGLINE": "ENERGY TO GO FURTHER"
        },
        "S229358414": {
            "ITEM_NAME": "KILIF ULTRA CONCENTRATED DETERGENT POWDER 2.5KG",
            "BARCODE": "5998866283771",
            "MANUFACTURER": "KILIF HOME CARE SDN BHD",
            "BRAND": "KILIF",
            "WEIGHT": "2.5KG",
            "PACKAGING_TYPE": "BAG",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "ULTRA POWER",
            "TYPE": "LAUNDRY DETERGENT",
            "FRAGRANCE_FLAVOR": "FLORAL FRESH",
            "PROMOTION": "",
            "ADDONS": "MEASURING SCOOP INCLUDED",
            "TAGLINE": "TOUGH ON STAINS"
        },
        "S229688224": {
            "ITEM_NAME": "KIVO CLASSIC TOMATO MIX SAUCE 60G",
            "BARCODE": "9557002019984",
            "MANUFACTURER": "KIVO FOOD INDUSTRIES",
            "BRAND": "KIVO",
            "WEIGHT": "60G",
            "PACKAGING_TYPE": "SACHET",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "CLASSIC",
            "TYPE": "TOMATO MIX / PASTE",
            "FRAGRANCE_FLAVOR": "RICH TOMATO",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": "RICH & THICK"
        },
        "S230256650": {
            "ITEM_NAME": "AJINOMOTO MONOSODIUM GLUTAMATE UMAMI SEASONING 200G",
            "BARCODE": "9556015001016",
            "MANUFACTURER": "AJINOMOTO (MALAYSIA) BERHAD",
            "BRAND": "AJINOMOTO",
            "WEIGHT": "200G",
            "PACKAGING_TYPE": "PACKET",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "STANDARD UMAMI",
            "TYPE": "FLAVOUR ENHANCER",
            "FRAGRANCE_FLAVOR": "SAVORY UMAMI",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": "EAT WELL LIVE WELL"
        },
        "S233065853": {
            "ITEM_NAME": "LIFE TOMATO KETCHUP SAUCE BOTTLE 485G",
            "BARCODE": "9556108001019",
            "MANUFACTURER": "YUM! BRANDS MALAYSIA / LIFE FOODS",
            "BRAND": "LIFE",
            "WEIGHT": "485G",
            "PACKAGING_TYPE": "BOTTLE",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "ORIGINAL",
            "TYPE": "SAUCE / CONDIMENT",
            "FRAGRANCE_FLAVOR": "SWEET & TANGY TOMATO",
            "PROMOTION": "NEW EASY POUR CAP",
            "ADDONS": "",
            "TAGLINE": "TASTE THE QUALITY"
        }
    }

    simulated_rows = []
    for pid, img_paths in groups.items():
        # Get mock extraction
        mock_raw = mock_catalog.get(pid, {
            "ITEM_NAME": f"GENERIC PRODUCT {pid}",
            "BARCODE": "1234567890123",
            "MANUFACTURER": "UNKNOWN",
            "BRAND": "UNKNOWN",
            "WEIGHT": "100G",
            "PACKAGING_TYPE": "BOX",
            "COUNTRY": "MALAYSIA",
            "VARIANT": "",
            "TYPE": "GROCERY",
            "FRAGRANCE_FLAVOR": "",
            "PROMOTION": "",
            "ADDONS": "",
            "TAGLINE": ""
        })
        
        # Simulate multi-image extraction with slight variations
        img_extractions = [mock_raw.copy()]
        if len(img_paths) > 1:
            variation = mock_raw.copy()
            variation["ITEM_NAME"] = variation["ITEM_NAME"].lower()
            img_extractions.append(variation)
        
        # Stage 3: Aggregate
        aggregated = pipeline.aggregate(img_extractions)
        # Stage 4: Validate and normalize
        validated = pipeline.validate_and_normalize(aggregated)
        # To IMDB row
        row = pipeline.to_imdb_row(validated)
        simulated_rows.append(row)

    output_file = PROJECT_ROOT / "IMDB_predictions.xlsx"
    pipeline.export_excel(simulated_rows, str(output_file))
    assert output_file.exists(), f"Failed to generate {output_file}"

    # Also test app.py's build_excel
    excel_buffer = app.build_excel(simulated_rows)
    assert excel_buffer.getbuffer().nbytes > 0, "app.build_excel returned empty buffer"

    # -------------------------------------------------------------
    # EXCEL VALIDATION & HACKATHON STYLING AUDIT
    # -------------------------------------------------------------
    print("\n--- [TEST 7] EXCEL FILE & STYLING COMPREHENSIVE AUDIT ---")
    wb = openpyxl.load_workbook(str(output_file))
    assert "IMDB" in wb.sheetnames, f"Sheet 'IMDB' missing, sheets found: {wb.sheetnames}"
    ws = wb["IMDB"]
    
    # Check dimensions
    max_r = ws.max_row
    max_c = ws.max_column
    assert max_r == 11, f"Expected 11 rows (1 header + 10 products), found {max_r}"
    assert max_c == 13, f"Expected 13 columns, found {max_c}"
    
    # Check headers
    read_headers = [ws.cell(row=1, column=c).value for c in range(1, 14)]
    assert read_headers == expected_cols, f"Header row mismatch:\nGot: {read_headers}\nExpected: {expected_cols}"
    
    # Check Header formatting (Calibri, Bold, 11pt, Center, Medium Border)
    for c in range(1, 14):
        cell = ws.cell(row=1, column=c)
        assert cell.font.bold is True, f"Header col {c} is not bold"
        assert cell.font.name == "Calibri", f"Header col {c} font name is {cell.font.name}"
        assert cell.font.size == 11, f"Header col {c} font size is {cell.font.size}"
        assert cell.alignment.horizontal == "center", f"Header col {c} align is {cell.alignment.horizontal}"
    
    # Check Freeze panes
    assert ws.freeze_panes == "A2", f"Freeze panes not set to A2: {ws.freeze_panes}"
    assert ws.row_dimensions[1].height == 20, f"Header row height is {ws.row_dimensions[1].height}"
    
    # Check Data rows formatting and content
    df_read = pd.read_excel(str(output_file), sheet_name="IMDB")
    print(f"\nRead {len(df_read)} rows via Pandas successfully:")
    print(df_read[["ITEM NAME", "BARCODE", "BRAND", "WEIGHT", "PACKAGING TYPE", "COUNTRY"]].to_string())
    
    # Verify all data values are uppercase and non-null for critical fields
    for idx, row in df_read.iterrows():
        item_name = str(row["ITEM NAME"])
        assert item_name == item_name.upper(), f"Row {idx} ITEM NAME not uppercase: {item_name}"
        barcode = str(row["BARCODE"])
        assert barcode.isdigit(), f"Row {idx} BARCODE not digits: {barcode}"
        brand = str(row["BRAND"])
        assert brand == brand.upper(), f"Row {idx} BRAND not uppercase: {brand}"
    
    test_results["excel_audit"] = {
        "status": "PASSED",
        "file_path": str(output_file),
        "rows": max_r - 1,
        "cols": max_c,
        "sheet_name": ws.title,
        "freeze_panes": ws.freeze_panes,
        "headers_verified": read_headers
    }

    print("\n" + "=" * 70)
    print("ALL TESTS PASSED SUCCESSFULLY (7/7 TEST SUITES)")
    print("=" * 70)
    
    return test_results

if __name__ == "__main__":
    results = run_tests()
    with open(PROJECT_ROOT / "test_report.json", "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
    print(f"\nSaved test report to {PROJECT_ROOT / 'test_report.json'}")
